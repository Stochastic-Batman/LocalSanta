import csv
import json
from openpyxl import load_workbook


def read_csv(filepath: str) -> dict[str, list[str]]:
    d: dict[str, list[str]] = {}
    with open(filepath, mode='r') as f:
        reader = csv.reader(f)
        for row in reader:
            d[row[0].strip()] = [x.strip() for x in row[1:]]
    return d


def read_excel(filepath: str) -> dict[str, list[str]]:
    d: dict[str, list[str]] = {}
    wb = load_workbook(filepath)
    sheet = wb.active
    if sheet is None:
        raise RuntimeError("No active sheet found in the workbook. When I had the same error, I realized I had saved it incorrectly, so I filled this Excel file in Excel Online and then downloaded it.")
    for row in sheet.iter_rows(values_only=True):
        if not row or all(cell is None for cell in row) or row[0] is None:
            continue
        d[row[0]] = [x.strip() for x in row[1:] if x is not None and x.strip() != ""]
    return d


def read_json(filepath: str) -> dict[str, list[str]]:
    d: dict[str, list[str]] = {}
    with open(filepath, "r", encoding="utf-8") as f:
        json_data = json.load(f)
    for k, v in json_data.items():
        d[k.strip()] = [x.strip() for x in v]
    return d